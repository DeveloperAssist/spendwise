"""Read a bank or UPI statement (CSV or Excel) into clean rows.

Real statements are messy:
- title lines above the table: "Account Statement", "Name: ..."; we find the header row ourselves;
- column names differ by bank: Narration, Particulars, Description, Paid to...;
- money out and money in come as one of three styles:
    1. separate Debit and Credit (or Withdrawal and Deposit) columns,
    2. one Amount column plus a Type column (DR / CR),
    3. one signed Amount column (negative = money out).

A broken row never stops an import: it's skipped and reported with its line number and reason.
Each row also gets a dedup key, so the same payment in two overlapping statements is stored once.
"""

import csv
import datetime as dt
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from functools import partial
from itertools import islice

from .clock import check_date
from .merchants import clean_merchant, merchant_key
from .money import is_negative, to_paise

DATE = ["date", "txn date", "transaction date", "value date", "tran date", "posting date"]
DESC = [
    "description",
    "narration",
    "particulars",
    "details",
    "remarks",
    "transaction details",
    "paid to",
    "merchant",
    "name",
    "transaction remarks",
]
AMOUNT = ["amount", "amount (inr)", "transaction amount", "amt"]
DEBIT = ["debit", "withdrawal", "withdrawal amt", "debit amount", "withdrawal amount", "dr amount"]
CREDIT = ["credit", "deposit", "deposit amt", "credit amount", "deposit amount", "cr amount"]
TYPE = ["type", "dr/cr", "cr/dr", "debit/credit", "transaction type", "txn type"]
TEXT_DATES = ["%Y-%m-%d", "%Y/%m/%d", "%d %b %Y", "%d-%b-%Y", "%d %B %Y", "%d-%b-%y", "%d %b %y"]
DAY_FIRST = ["%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%y", "%d/%m/%y"]
MONTH_FIRST = ["%m-%d-%Y", "%m/%d/%Y", "%m.%d.%Y", "%m-%d-%y", "%m/%d/%y"]
NUMERIC_DATE = re.compile(r"^\s*(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})\s*$")
# balance lines are not payments. Only exact markers: "TOTAL MALL" or "Total Fitness" ARE payments.
SKIP_EXACT = {"total", "grand total", "b/f", "c/f"}
SKIP_PREFIX = ("opening balance", "closing balance", "balance b/f", "balance c/f", "balance forward")
MAX_COLS = 40
MAX_UNZIPPED = 20 * 1024 * 1024


@dataclass
class ParsedRow:
    date: dt.date
    description: str
    merchant: str
    amount_paise: int
    kind: str  # debit | credit
    dedup_key: str = ""


@dataclass
class ParseResult:
    rows: list[ParsedRow] = field(default_factory=list)
    skipped: list[dict] = field(default_factory=list)  # {"line": 7, "reason": "..."}


def day_first(values) -> bool:
    """Is 01/02/2026 the 1st of February (India) or January 2nd (US)? A file can't mix the two, so
    any date with a first part above 12 settles it as day-first, a second part above 12 as month-first.
    If nothing settles it, assume day-first, like Indian banks."""
    for v in values:
        m = NUMERIC_DATE.match(str(v or ""))
        if m:
            a, b = int(m[1]), int(m[2])
            if a > 12:
                return True
            if b > 12:
                return False
    return True


def parse_date(value, dayfirst: bool = True) -> dt.date:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value or "").strip()
    if len(text) > 10 and text[10:11] in (" ", "T"):  # "2026-09-30 00:00:00"
        text = text[:10]
    for f in TEXT_DATES + (DAY_FIRST if dayfirst else MONTH_FIRST):
        try:
            return dt.datetime.strptime(text, f).date()
        except ValueError:
            pass
    raise ValueError(f"unknown date {text[:30]!r}")


def _norm(h) -> str:
    return " ".join(str(h or "").replace("﻿", "").strip().lower().replace(".", "").split())


def _find(headers: list[str], names: list[str]) -> int | None:
    for n in names:
        if n in headers:
            return headers.index(n)
    return None


def _header_index(table: list[list]) -> int:
    """The first row (of the first 25) that has a date column and some kind of amount column."""
    for i, row in enumerate(table[:25]):
        h = [_norm(c) for c in row]
        if _find(h, DATE) is not None and any(_find(h, x) is not None for x in (AMOUNT, DEBIT, CREDIT)):
            return i
    raise ValueError("no header row found: the file needs a Date column and an Amount (or Debit/Credit) column")


def _read_xlsx(data: bytes, limit: int) -> list[list]:
    # An .xlsx is a zip. Check the unzipped size first, and never trust the sheet's own size claim:
    # a 5 KB file can declare 1,048,576 rows x 16,384 columns, and reading that padded grid would eat
    # all the server's memory (found in review).
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED:
                raise ValueError("the Excel file is too large once unzipped")
    except zipfile.BadZipFile:
        raise ValueError("could not read the Excel file: it isn't a valid .xlsx") from None
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            ws.reset_dimensions()
            return [list(r) for r in islice(ws.iter_rows(values_only=True, max_col=MAX_COLS), limit)]
        finally:
            wb.close()
    except ValueError:
        raise
    except Exception as e:  # a corrupt sheet fails while reading, not only while opening
        raise ValueError(f"could not read the Excel file: {e.__class__.__name__}") from None


def read_table(data: bytes, filename: str, limit: int = 20_031) -> list[list]:
    """CSV or XLSX bytes -> at most `limit` rows of at most MAX_COLS cells."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(data, limit)
    if name.endswith(".csv") or "." not in name:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252", errors="replace")  # old Excel CSV exports
        try:
            return [row[:MAX_COLS] for row in islice(csv.reader(io.StringIO(text)), limit)]
        except csv.Error as e:  # e.g. a field over the size limit
            raise ValueError(f"could not read the CSV file: {e}") from None
    raise ValueError("unsupported file type: upload a .csv or .xlsx statement")


def parse_statement(data: bytes, filename: str, max_rows: int = 20_000) -> ParseResult:
    table = read_table(data, filename, limit=max_rows + 31)  # header lines + one extra row to spot "too many"
    hi = _header_index(table)
    h = [_norm(c) for c in table[hi]]
    col_date, col_desc = _find(h, DATE), _find(h, DESC)
    col_amt, col_dr, col_cr, col_type = _find(h, AMOUNT), _find(h, DEBIT), _find(h, CREDIT), _find(h, TYPE)
    if col_desc is None:
        raise ValueError("no description column (looked for: Narration, Description, Particulars, Paid to...)")
    body = table[hi + 1 :]
    if len(body) > max_rows:
        raise ValueError(f"too many rows ({len(body)}); the limit is {max_rows}")

    # style 3 needs to know whether the file uses signs at all
    signed = (
        col_amt is not None
        and col_type is None
        and col_dr is None
        and any(is_negative(r[col_amt]) for r in body if col_amt < len(r) and r[col_amt] not in (None, ""))
    )

    dayfirst = day_first(_cell(r, col_date) for r in body)
    out = ParseResult()
    for line, row in enumerate(body, start=hi + 2):  # spreadsheet line numbers (1-based)
        cell = partial(_cell, row)
        if not any(str(c or "").strip() for c in row):
            continue  # blank line
        desc = " ".join(str(cell(col_desc)).split())
        if desc.lower() in SKIP_EXACT or desc.lower().startswith(SKIP_PREFIX):
            continue
        try:
            date = check_date(parse_date(cell(col_date), dayfirst))
            if not desc:
                raise ValueError("no description")
            if col_dr is not None or col_cr is not None:  # style 1
                dr, cr = str(cell(col_dr)).strip(), str(cell(col_cr)).strip()
                if dr not in ("", "0", "0.0", "0.00", "-"):
                    amount, kind = to_paise(dr), "debit"
                elif cr not in ("", "0", "0.0", "0.00", "-"):
                    amount, kind = to_paise(cr), "credit"
                else:
                    raise ValueError("no debit or credit amount")
            else:
                raw = cell(col_amt)
                amount = to_paise(raw)
                if col_type is not None:  # style 2
                    t = str(cell(col_type)).strip().lower().rstrip(".")
                    # exact short codes first: "d" must not catch "deposit"
                    if t in ("d", "dr") or t.startswith(("debit", "withdraw")):
                        kind = "debit"
                    elif t in ("c", "cr") or t.startswith(("credit", "deposit")):
                        kind = "credit"
                    else:
                        raise ValueError(f"unknown type {t!r}")
                else:  # style 3 (or a payments-only export)
                    kind = "debit" if (is_negative(raw) or not signed) else "credit"
                    if str(raw).strip().upper().endswith("CR"):
                        kind = "credit"
        except ValueError as e:
            out.skipped.append({"line": line, "reason": str(e)})
            continue
        out.rows.append(
            ParsedRow(date=date, description=desc[:300], merchant=clean_merchant(desc), amount_paise=amount, kind=kind)
        )
    _add_dedup_keys(out.rows)
    return out


def _cell(row: list, i: int | None):
    return row[i] if i is not None and i < len(row) and row[i] is not None else ""


def _add_dedup_keys(rows: list[ParsedRow]) -> None:
    """Same date + merchant + amount + kind -> same key. Two genuine identical payments on one day
    (two chais, ₹20 each) get #0 and #1, so both are kept, and re-importing still matches them."""
    seen: dict[tuple, int] = {}
    for r in rows:
        base = (r.date.isoformat(), merchant_key(r.merchant), r.amount_paise, r.kind)
        n = seen.get(base, 0)
        seen[base] = n + 1
        r.dedup_key = hashlib.sha256("|".join(map(str, (*base, n))).encode()).hexdigest()[:40]
